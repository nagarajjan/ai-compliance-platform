"""
Creates sample financial source Excel file: financial_audit_metrics.xlsx
Contains:
1. Quarterly_Financial_Performance
2. Regional_Financial_Distribution
3. Compliance_Security_Budget
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

def create_financial_excel():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font   = Font(name="Calibri", size=11)
    bold_font   = Font(name="Calibri", size=11, bold=True)
    align_left  = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_center= Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # ── SHEET 1: Quarterly_Financial_Performance ─────────────────────────
    ws1 = wb.create_sheet(title="Quarterly_Financial_Performance")
    headers1 = ["Quarter", "Gross_Revenue_USD_M", "Operating_Expenses_USD_M", "EBITDA_USD_M", "Net_Profit_USD_M", "Net_Margin_Pct"]
    ws1.append(headers1)

    data1 = [
        ["Q1-2026", 128.5, 74.2, 42.1, 32.4, 25.2],
        ["Q2-2026", 142.8, 81.0, 48.5, 37.8, 26.5],
        ["Q3-2026", 156.2, 88.6, 52.3, 41.2, 26.4],
        ["Q4-2026 (Est)", 175.0, 96.5, 61.0, 48.5, 27.7],
    ]
    for row in data1:
        ws1.append(row)

    # ── SHEET 2: Regional_Financial_Distribution ─────────────────────────
    ws2 = wb.create_sheet(title="Regional_Financial_Distribution")
    headers2 = ["Region", "Revenue_USD_M", "Profit_USD_M", "Growth_YoY_Pct", "Target_Revenue_USD_M", "Variance_Pct", "Status"]
    ws2.append(headers2)

    data2 = [
        ["North America", 245.0, 68.5, 14.2, 230.0, 6.5, "Outperforming"],
        ["EMEA", 165.2, 41.8, 8.5, 170.0, -2.8, "On Target"],
        ["APAC", 132.8, 38.4, 22.4, 120.0, 10.7, "High Growth"],
        ["LATAM", 42.5, 9.2, 11.0, 45.0, -5.6, "Under Target"],
        ["Middle East", 17.0, 2.1, -4.2, 25.0, -32.0, "Critical Lag"],
    ]
    for row in data2:
        ws2.append(row)

    # ── SHEET 3: Compliance_Security_Budget ───────────────────────────────
    ws3 = wb.create_sheet(title="Compliance_Security_Budget")
    headers3 = ["Category", "Allocated_Budget_kUSD", "Actual_Spend_kUSD", "Variance_kUSD", "Variance_Pct", "Budget_Status"]
    ws3.append(headers3)

    data3 = [
        ["Cloud Infrastructure & ZTNA", 1450.0, 1620.0, 170.0, 11.7, "Over Budget"],
        ["AML & Sanctions Screening", 850.0, 810.0, -40.0, -4.7, "Under Budget"],
        ["Penetration Testing & Remediation", 600.0, 780.0, 180.0, 30.0, "Critical Over-Run"],
        ["Cryptographic Key Mgmt (HSM)", 450.0, 430.0, -20.0, -4.4, "On Budget"],
        ["Regulatory Audits (MAS/ISO)", 350.0, 345.0, -5.0, -1.4, "On Budget"],
        ["Data Loss Prevention (DLP)", 500.0, 490.0, -10.0, -2.0, "On Budget"],
    ]
    for row in data3:
        ws3.append(row)

    # Apply formatting across all sheets
    for ws in [ws1, ws2, ws3]:
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center

        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = data_font
                cell.border = thin_border
                val = cell.value
                if isinstance(val, (int, float)):
                    cell.alignment = align_right
                else:
                    cell.alignment = align_left

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Save to both locations
    p1 = Path("vectorstore/demo/raw/financial_audit_metrics.xlsx")
    p2 = Path("sample_documents/financial_audit_metrics.xlsx")
    p1.parent.mkdir(parents=True, exist_ok=True)
    p2.parent.mkdir(parents=True, exist_ok=True)

    wb.save(str(p1))
    wb.save(str(p2))
    print(f"Created: {p1.resolve()}")
    print(f"Created: {p2.resolve()}")

if __name__ == "__main__":
    create_financial_excel()
